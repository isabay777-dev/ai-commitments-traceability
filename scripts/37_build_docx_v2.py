#!/usr/bin/env python3
"""Run with the bundled documents Python runtime. Pandoc conversion + layout only.

Templates are converted from article1 with the requested pandoc command.
The post-pass changes styles, widths and pagination; never manuscript wording.
"""
from pathlib import Path
import subprocess
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
ART=ROOT/'article1'
files=sorted(ART.glob('ЧЕРНОВИК_RU_*.md'))+[ART/'tables_v2/Tables_I_IV.md', ART/'tables_v2/Figure_1.md']
if len(sys.argv)>1:
    files=[p for p in files if p.name in sys.argv[1:]]
for md in files:
    output=md.with_suffix('.docx')
    args=['pandoc','-f','markdown-yaml_metadata_block',str(md.relative_to(ART)),'-o',str(output.relative_to(ART))]
    if md.name=='Figure_1.md': args+=['--resource-path=tables_v2']
    subprocess.run(args,cwd=ART,check=True)
    doc=Document(output)
    for section in doc.sections:
        section.page_width=Inches(8.27);section.page_height=Inches(11.69)
        section.top_margin=Inches(.7);section.bottom_margin=Inches(.7)
        section.left_margin=Inches(.65);section.right_margin=Inches(.65)
    for style in doc.styles:
        if style.type==1:
            style.font.name='Arial'
            style.font.color.rgb=RGBColor(0,0,0)
            color=style.element.find('.//'+qn('w:color'))
            if color is not None:
                for attr in ['themeColor','themeShade','themeTint']:
                    color.attrib.pop(qn('w:'+attr),None)
    for name in ['Normal','Body Text','First Paragraph']:
        doc.styles[name].font.size=Pt(10.5)
        doc.styles[name].paragraph_format.space_after=Pt(6)
        doc.styles[name].paragraph_format.line_spacing=1.08
        if md.name=='ЧЕРНОВИК_RU_3_Данные_и_методы.md':
            doc.styles[name].font.size=Pt(10.5)
            doc.styles[name].paragraph_format.space_after=Pt(5)
            doc.styles[name].paragraph_format.line_spacing=1.02
    for name,size in [('Title',19),('Heading 1',17),('Heading 2',14),('Heading 3',12),('Heading 4',11)]:
        doc.styles[name].font.size=Pt(size)
        doc.styles[name].paragraph_format.keep_with_next=True
    # The first Markdown heading is the document title, not a numbered section.
    doc.paragraphs[0].style=doc.styles['Title']
    for para in doc.paragraphs:
        if para.text.startswith(('[Table ', '[Figure ')):
            para.paragraph_format.keep_with_next=True
        if para.text.startswith('Table ') and md.name=='Tables_I_IV.md':
            para.paragraph_format.page_break_before=not para.text.startswith('Table I ·')
        for run in para.runs:
            if para.style.name in ['Title','Subtitle'] or para.style.name.startswith('Heading'):
                run.font.color.rgb=RGBColor(0,0,0)
    for table in doc.tables:
        table.autofit=False
        n=len(table.columns)
        header=[c.text for c in table.rows[0].cells]
        if md.name=='Tables_I_IV.md':
            if header[0]=='Панель': widths=[.65,1.92,1.10,1.10,1.10,1.10]
            elif header[0]=='Версия': widths=[1.45,.68,.38,1.40,1.53,1.53]
            elif header[0]=='Оценка': widths=[2.0,1.45,.7,2.82]
            elif header[0]=='Страна / сектор': widths=[1.45,1.65,1.65,1.5,.72]
            elif header[0]=='Страна': widths=[.70,1.05,1.00,1.52,1.25,1.45]
            else: widths=[2.05,1.80,2.12,1.00]
        elif header[0]=='Код': widths=[1.40,5.57]
        elif header[0]=='Стадия': widths=[3.37,.90,.90,.90,.90]
        elif header[0]=='Вопрос': widths=[2.35,3.20,1.42]
        elif n==2: widths=[2.05,4.92]
        else: widths=[6.97/n]*n
        assert len(widths)==n
        for col,width in zip(table.columns,widths): col.width=Inches(width)
        for rownum,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr()
            avoid=OxmlElement('w:cantSplit');trpr.append(avoid)
            if rownum==0:
                repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for cell,width in zip(row.cells,widths):
                cell.width=Inches(width)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(4)
                    para.paragraph_format.space_before=Pt(4)
                    para.paragraph_format.line_spacing=1.0
                    for run in para.runs:
                        run.font.name='Arial';run.font.size=Pt(9.5)
                        if rownum==0: run.bold=True
    doc.core_properties.author=''
    doc.core_properties.last_modified_by=''
    doc.save(output)
    print(output.relative_to(ROOT))
